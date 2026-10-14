"""Fig. 1, Fig. 2 를 손으로 고칠 수 있는 pptx 로 만든다 (모든 상자, 글자, 화살표가 개별 도형).

슬라이드 크기 = 논문 그림 크기. PowerPoint 에서 고친 뒤 '내보내기 > PDF' 로 저장하면
figures/ 의 PDF 를 그대로 교체할 수 있다. 좌표는 fig_framework.py, fig_vessel.py 와 같다.
"""
import os
import copy

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
OUTDIR = os.path.join(HERE, '..', 'ppt')
PH = os.path.join(HERE, 'photos')
os.makedirs(OUTDIR, exist_ok=True)

FONT = 'Times New Roman'
MATHFONT = 'Cambria Math'

C_PLANT, C_PLANT_BG = '2E5A87', 'E8EFF6'
C_OFF, C_OFF_BG = 'A65E1A', 'FBEFE3'
C_ON, C_ON_BG = '2F7D4F', 'E7F3EC'
C_GRAY, C_GRAY_BG = '6B6B6B', 'F2F2F2'
C_RED = '9B2222'
C_YEL = 'FFD000'


# ---------- 글자 서식 ----------
# 한 줄은 조각(run) 목록. 조각 = (글자, 서식문자열). 서식: i 기울임, b 굵게, s 아래첨자, m 수식글꼴
def T(*parts):
    out = []
    for p in parts:
        out.append((p, '') if isinstance(p, str) else p)
    return out


def I(t):
    return (t, 'i')


def SUB(t):
    return (t, 's')


def ISUB(t):
    return (t, 'is')


def M(t):
    return (t, 'm')


def fill_text(tf, lines, size, color='000000', align=PP_ALIGN.CENTER, bold=False, italic=False,
              anchor=MSO_ANCHOR.MIDDLE, spacing=1.0):
    tf.clear()
    tf.word_wrap = False
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for li, line in enumerate(lines):
        para = tf.paragraphs[0] if li == 0 else tf.add_paragraph()
        para.alignment = align
        para.line_spacing = spacing
        for text, fmt in line:
            r = para.add_run()
            r.text = text
            f = r.font
            f.size = Pt(size)
            f.name = MATHFONT if 'm' in fmt else FONT
            f.bold = bold or ('b' in fmt)
            f.italic = italic or ('i' in fmt)
            f.color.rgb = RGBColor.from_string(color)
            if 's' in fmt:
                r.font._rPr.set('baseline', '-25000')


# ---------- 도형 ----------
class Canvas:
    """matplotlib 와 같은 좌표계 (인치, y 는 아래에서 위로) 를 pptx 좌표로 바꿔 그린다."""

    def __init__(self, slide, H):
        self.s = slide
        self.H = H

    def Y(self, y):  # 위쪽 기준 좌표
        return self.H - y

    def rbox(self, x, y0, w, y1, ec, fc='FFFFFF', lw=0.7, radius=0.04, dash=False, name=None):
        h = y1 - y0
        sh = self.s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(self.Y(y1)),
                                     Inches(w), Inches(h))
        sh.adjustments[0] = min(0.5, radius / min(w, h))
        sh.fill.solid()
        sh.fill.fore_color.rgb = RGBColor.from_string(fc)
        sh.line.color.rgb = RGBColor.from_string(ec)
        sh.line.width = Pt(lw)
        if dash:
            sh.line.dash_style = 4
        sh.shadow.inherit = False
        if name:
            sh.name = name
        return sh

    def box(self, x, y0, w, y1, lines, ec, fc='FFFFFF', size=7.0, color='000000', italic=False,
            name=None, lw=0.7):
        sh = self.rbox(x, y0, w, y1, ec, fc, lw=lw, name=name)
        fill_text(sh.text_frame, lines, size, color=color, italic=italic, spacing=1.05)
        return sh

    def text(self, x, y, lines, size=6.5, color='000000', ha='center', va='center', bold=False,
             italic=False, w=None, h=None, name=None):
        nlines = len(lines)
        h = h or (size / 72.0) * 1.25 * nlines + 0.02
        w = w or 2.5
        if ha == 'center':
            left = x - w / 2
        elif ha == 'left':
            left = x
        else:
            left = x - w
        if va == 'center':
            top = self.Y(y) - h / 2
        elif va == 'bottom':
            top = self.Y(y) - h
        else:
            top = self.Y(y)
        tb = self.s.shapes.add_textbox(Inches(left), Inches(top), Inches(w), Inches(h))
        align = {'center': PP_ALIGN.CENTER, 'left': PP_ALIGN.LEFT, 'right': PP_ALIGN.RIGHT}[ha]
        fill_text(tb.text_frame, lines, size, color=color, align=align, bold=bold, italic=italic,
                  spacing=1.05)
        if name:
            tb.name = name
        return tb

    def _arrowhead(self, shape, color):
        ln = shape.line._get_or_add_ln()
        tail = etree.SubElement(ln, qn('a:tailEnd'))
        tail.set('type', 'triangle')
        tail.set('w', 'sm')
        tail.set('len', 'sm')

    def arrow(self, p0, p1, color='000000', lw=0.9, head=True):
        c = self.s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(p0[0]), Inches(self.Y(p0[1])),
                                        Inches(p1[0]), Inches(self.Y(p1[1])))
        c.line.color.rgb = RGBColor.from_string(color)
        c.line.width = Pt(lw)
        if head:
            self._arrowhead(c, color)
        return c

    def poly(self, pts, color='000000', lw=0.9, head=True):
        """꺾은선 화살표 (마지막 점에 화살촉)."""
        e = [(Inches(x), Inches(self.Y(y))) for x, y in pts]
        fb = self.s.shapes.build_freeform(e[0][0], e[0][1], scale=1.0)
        fb.add_line_segments(e[1:], close=False)
        sh = fb.convert_to_shape()
        sh.fill.background()
        sh.line.color.rgb = RGBColor.from_string(color)
        sh.line.width = Pt(lw)
        sh.shadow.inherit = False
        if head:
            self._arrowhead(sh, color)
        return sh

    def ellipse(self, cx, cy, a, b, color, alpha):
        sh = self.s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - a / 2), Inches(self.Y(cy) - b / 2),
                                     Inches(a), Inches(b))
        sh.fill.solid()
        sh.fill.fore_color.rgb = RGBColor.from_string(color)
        clr = sh.fill._xPr.find(qn('a:solidFill')).find(qn('a:srgbClr'))
        al = etree.SubElement(clr, qn('a:alpha'))
        al.set('val', str(int(alpha * 100000)))
        sh.line.color.rgb = RGBColor.from_string(color)
        sh.line.width = Pt(0.6)
        sh.shadow.inherit = False
        return sh


class TC:
    """위쪽 왼쪽이 원점인 좌표계 (인치) 로 그리는 도구. Fig. 1 용."""

    def __init__(self, slide):
        self.s = slide

    def rect(self, x, y, w, h, ec, fc='FFFFFF', lw=0.7, radius=0.04, dash=False, name=None):
        sh = self.s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        sh.adjustments[0] = min(0.5, radius / min(w, h))
        if fc is None:
            sh.fill.background()
        else:
            sh.fill.solid()
            sh.fill.fore_color.rgb = RGBColor.from_string(fc)
        sh.line.color.rgb = RGBColor.from_string(ec)
        sh.line.width = Pt(lw)
        if dash:
            sh.line.dash_style = 4
        sh.shadow.inherit = False
        if name:
            sh.name = name
        return sh

    def node(self, x, y, w, h, lines, ec, fc='FFFFFF', size=7.0, color='000000', lw=0.7, bold=False, name=None):
        sh = self.rect(x, y, w, h, ec, fc, lw=lw, name=name)
        fill_text(sh.text_frame, lines, size, color=color, bold=bold, spacing=1.05)
        return sh

    def label(self, x, y, w, h, lines, size=6.5, color='000000', align='center', bold=False, italic=False,
              anchor=MSO_ANCHOR.MIDDLE, name=None):
        tb = self.s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        al = {'center': PP_ALIGN.CENTER, 'left': PP_ALIGN.LEFT, 'right': PP_ALIGN.RIGHT}[align]
        fill_text(tb.text_frame, lines, size, color=color, align=al, bold=bold, italic=italic,
                  anchor=anchor, spacing=1.05)
        if name:
            tb.name = name
        return tb

    def _end(self, shape, tag):
        ln = shape.line._get_or_add_ln()
        e = etree.SubElement(ln, qn(tag))
        e.set('type', 'triangle')
        e.set('w', 'sm')
        e.set('len', 'sm')

    def arrow(self, p0, p1, color='000000', lw=0.9, both=False, head=True, dash=False):
        c = self.s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(p0[0]), Inches(p0[1]),
                                        Inches(p1[0]), Inches(p1[1]))
        c.line.color.rgb = RGBColor.from_string(color)
        c.line.width = Pt(lw)
        if dash:
            c.line.dash_style = 4
        if both:
            self._end(c, 'a:headEnd')
        if head:
            self._end(c, 'a:tailEnd')
        return c

    def poly(self, pts, color='000000', lw=0.9, head=True):
        e = [(Inches(x), Inches(y)) for x, y in pts]
        fb = self.s.shapes.build_freeform(e[0][0], e[0][1], scale=1.0)
        fb.add_line_segments(e[1:], close=False)
        sh = fb.convert_to_shape()
        sh.fill.background()
        sh.line.color.rgb = RGBColor.from_string(color)
        sh.line.width = Pt(lw)
        sh.shadow.inherit = False
        if head:
            self._end(sh, 'a:tailEnd')
        return sh

    def polygon(self, pts, fc, ec=None, lw=0.6, name=None):
        e = [(Inches(x), Inches(y)) for x, y in pts]
        fb = self.s.shapes.build_freeform(e[0][0], e[0][1], scale=1.0)
        fb.add_line_segments(e[1:], close=True)
        sh = fb.convert_to_shape()
        sh.fill.solid()
        sh.fill.fore_color.rgb = RGBColor.from_string(fc)
        if ec:
            sh.line.color.rgb = RGBColor.from_string(ec)
            sh.line.width = Pt(lw)
        else:
            sh.line.fill.background()
        sh.shadow.inherit = False
        if name:
            sh.name = name
        return sh

    def ellipse(self, cx, cy, a, b, color, alpha):
        sh = self.s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - a / 2), Inches(cy - b / 2), Inches(a), Inches(b))
        sh.fill.solid()
        sh.fill.fore_color.rgb = RGBColor.from_string(color)
        clr = sh.fill._xPr.find(qn('a:solidFill')).find(qn('a:srgbClr'))
        etree.SubElement(clr, qn('a:alpha')).set('val', str(int(alpha * 100000)))
        sh.line.color.rgb = RGBColor.from_string(color)
        sh.line.width = Pt(0.6)
        sh.shadow.inherit = False
        return sh

    def boat(self, x, y, L, B, color=C_PLANT):
        """위에서 본 선체 (선미 왼쪽, 선수 오른쪽) 와 선미 드라이브."""
        hull = [(x, y + 0.12 * B), (x + 0.62 * L, y), (x + 0.86 * L, y + 0.12 * B), (x + L, y + 0.5 * B),
                (x + 0.86 * L, y + 0.88 * B), (x + 0.62 * L, y + B), (x, y + 0.88 * B)]
        self.polygon(hull, color, name='boat_hull')
        self.polygon([(x - 0.07 * L, y + 0.38 * B), (x, y + 0.38 * B), (x, y + 0.62 * B),
                      (x - 0.07 * L, y + 0.62 * B)], '1B3A5C', name='boat_sterndrive')

    def gusts(self, x, y, n=3, dy=0.07, L=0.16, color='8A8A8A'):
        for i in range(n):
            self.arrow((x, y + i * dy), (x + L, y + i * dy), color=color, lw=0.8)


def new_pres(W, H):
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # 빈 레이아웃
    return prs, slide


# ======================= Fig. 1 =======================
def make_fig1():
    """왼쪽 Offline 학습(시뮬레이션), 오른쪽 Online 배포(실선). 학습된 네트워크는 주황 테두리."""
    W, H = 7.16, 2.45
    prs, slide = new_pres(W, H)
    t = TC(slide)

    # ----- 두 영역 -----
    OX, OY, OW, OH = 0.04, 0.04, 4.38, 2.37   # Offline
    NX, NY, NW, NH = 4.86, 0.04, 2.26, 2.37   # Online
    t.rect(OX, OY, OW, OH, C_OFF, C_OFF_BG, lw=1.0, radius=0.08, name='region_offline')
    t.rect(NX, NY, NW, NH, C_ON, C_ON_BG, lw=1.0, radius=0.08, name='region_online')
    t.label(OX + 0.10, OY + 0.05, 3.0, 0.22, [T(('Offline', 'b'), ' training in simulation')],
            size=8.0, color=C_OFF, align='left')
    t.label(NX + 0.10, NY + 0.05, 2.1, 0.22, [T(('Online', 'b'), ' deployment on the vessel')],
            size=8.0, color=C_ON, align='left')

    # ----- Offline 안쪽 -----
    # 시뮬레이터
    sx, sy, sw, sh = OX + 0.12, 0.40, 1.02, 1.88
    t.rect(sx, sy, sw, sh, C_PLANT, 'FFFFFF', name='simulator')
    t.label(sx, sy + 0.06, sw, 0.18, [T('Simulator')], size=7.2, bold=True, color=C_PLANT)
    t.boat(sx + 0.20, sy + 0.74, 0.68, 0.25)
    t.gusts(sx + 0.10, sy + 0.42, n=3, dy=0.06, L=0.15)
    t.label(sx + 0.30, sy + 0.41, 0.6, 0.14, [T(I('d'), ' ∈ ', M('𝒲'))], size=6.5, color=C_GRAY, align='left')
    t.label(sx, sy + 1.18, sw, 0.62, [T('gear ', I('g'), ' ∈ {R, N, F}'),
                                      T('dwell ', I('D'), ISUB('g'), ' = 3 s'),
                                      T('steering ', I('δ'))], size=6.5)

    # Stage 1
    s1x, s1y, s1w, s1h = 1.42, 0.40, 2.86, 0.98
    t.rect(s1x, s1y, s1w, s1h, C_OFF, 'FFFFFF', name='stage1')
    t.label(s1x + 0.08, s1y + 0.04, 2.6, 0.18, [T(('Stage 1', 'b'), '  dwell indexed safety value')],
            size=7.0, color=C_OFF, align='left')
    nh = 0.26
    t.node(s1x + 0.12, s1y + 0.30, 0.46, nh, [T(I('π'), ISUB('h'))], C_OFF, size=7.5)
    t.node(s1x + 0.12, s1y + 0.64, 0.46, nh, [T(I('μ'))], C_OFF, size=7.5)
    t.node(s1x + 0.98, s1y + 0.40, 0.52, 0.40, [T(I('Q'), ISUB('h'))], C_OFF, size=7.5)
    t.arrow((s1x + 0.58, s1y + 0.43), (s1x + 0.98, s1y + 0.53), color=C_OFF)
    t.arrow((s1x + 0.58, s1y + 0.77), (s1x + 0.98, s1y + 0.67), color=C_OFF)
    t.label(s1x + 0.60, s1y + 0.27, 0.36, 0.14, [T('max')], size=6.2, color=C_GRAY)
    t.label(s1x + 0.60, s1y + 0.80, 0.36, 0.14, [T('min')], size=6.2, color=C_GRAY)
    # 안전집합 족
    ex, ey = s1x + 2.20, s1y + 0.56
    for a, b, al in [(0.86, 0.48, 0.16), (0.58, 0.32, 0.30), (0.30, 0.17, 0.55)]:
        t.ellipse(ex, ey, a, b, C_OFF, al)
    t.label(ex - 0.43, s1y + 0.13, 0.86, 0.14, [T(M('𝒮'), '(', I('g'), ', ', I('τ'), ISUB('g'), ')')],
            size=6.5, color=C_OFF)
    for dx, lab, col in [(0.36, '0', C_OFF), (0.22, '3', C_OFF), (0.0, '5', 'FFFFFF')]:
        t.label(ex + dx - 0.05, ey - 0.07, 0.10, 0.14, [T(lab)], size=6.0, color=col)
    t.arrow((s1x + 1.50, s1y + 0.60), (ex - 0.45, ey + 0.04), color=C_OFF)

    # Stage 2
    s2x, s2y, s2w, s2h = 1.42, 1.48, 2.86, 0.80
    t.rect(s2x, s2y, s2w, s2h, C_OFF, 'FFFFFF', name='stage2')
    t.label(s2x + 0.08, s2y + 0.04, 2.6, 0.18, [T(('Stage 2', 'b'), '  task policy')],
            size=7.0, color=C_OFF, align='left')
    t.node(s2x + 0.12, s2y + 0.34, 0.86, 0.32, [T(I('Q'), ISUB('h'), ' frozen')], C_GRAY, C_GRAY_BG, size=7.0)
    t.node(s2x + 1.40, s2y + 0.34, 0.56, 0.32, [T(I('π'), ISUB('t'))], C_OFF, size=7.5)
    t.arrow((s2x + 0.98, s2y + 0.50), (s2x + 1.40, s2y + 0.50), color=C_OFF)
    t.label(s2x + 2.02, s2y + 0.30, 0.80, 0.40, [T('reward'), T('s.t. ', I('q'), ' ≥ ', I('τ'))], size=6.5)
    # Stage 1 에서 Stage 2 로
    t.arrow((s1x + 0.40, s1y + s1h), (s2x + 0.40, s2y + 0.34), color=C_OFF)

    # 시뮬레이터 왕복 (rollout)
    for yy in (s1y + s1h / 2, s2y + s2h / 2):
        t.arrow((sx + sw, yy), (s1x, yy), color=C_PLANT, both=True)

    # ----- Offline 에서 Online 으로 -----
    ym = 1.24
    t.arrow((OX + OW, ym), (NX, ym), color='000000', lw=1.4)
    t.label(OX + OW - 0.02, ym - 0.24, 0.48, 0.16, [T('deploy')], size=6.5)
    t.label(OX + OW - 0.02, ym + 0.06, 0.48, 0.16, [T('frozen')], size=6.5, color=C_GRAY)

    # ----- Online 안쪽 (폐루프) -----
    t.node(NX + 0.14, 0.40, 0.56, 0.32, [T('EKF')], C_ON, size=7.0)
    t.node(NX + 0.86, 0.40, 0.50, 0.32, [T(I('π'), ISUB('t'))], C_OFF, size=7.5)
    t.node(NX + 1.50, 0.40, 0.64, 0.46, [T('Shield'), T(I('q'), ' ≥ ', I('τ')), T('else ', I('π'), ISUB('h'))], C_OFF, size=6.8)
    t.arrow((NX + 0.70, 0.56), (NX + 0.86, 0.56), color=C_ON)
    t.arrow((NX + 1.36, 0.56), (NX + 1.50, 0.56), color=C_ON)
    # 선박
    vx, vy, vw, vh = NX + 0.30, 1.46, 1.62, 0.82
    t.rect(vx, vy, vw, vh, C_PLANT, 'FFFFFF', name='vessel')
    t.boat(vx + 0.42, vy + 0.16, 0.86, 0.30)
    t.gusts(vx + 0.10, vy + 0.08, n=3, dy=0.07, L=0.18)
    t.label(vx, vy + 0.52, vw, 0.24, [T('Vessel, PLC, servo cylinders')], size=6.5)
    # 명령과 측정
    t.poly([(NX + 1.82, 0.86), (NX + 1.82, vy)], color=C_ON)
    t.label(NX + 1.86, 1.12, 0.40, 0.24, [T(I('g'), ', ', I('δ'))], size=6.8, align='left')
    t.poly([(NX + 0.42, vy), (NX + 0.42, 0.72)], color=C_ON)
    t.label(NX + 0.02, 1.00, 0.38, 0.34, [T('GNSS'), T('IMU')], size=6.3, color=C_GRAY, align='right')
    t.label(NX + 0.62, 1.10, 1.10, 0.16, [T('every 0.5 s')], size=6.3, color=C_ON)

    out = os.path.join(OUTDIR, 'fig1_framework.pptx')
    prs.save(out)
    return out


# ======================= Fig. 2 =======================
C_LINK = '00A0E0'   # 연결선 (원본 5번 슬라이드의 하늘색)
C_PROTO = '1F6FB2'  # 통신 방식 글자
C_BOUND = 'D62828'  # 선박 제어부 점선


def _media(name):
    """원본 'Ship mechanical system.pptx' 의 그림 파일을 읽는다."""
    import io
    import zipfile
    z = zipfile.ZipFile(os.path.join(HERE, '..', '..', '..', 'Ship mechanical system.pptx'))
    return Image.open(io.BytesIO(z.read('ppt/media/' + name))).convert('RGB')


def _crop_frac(im, l, t, r, b):
    w, h = im.size
    return im.crop((int(l * w), int(t * h), int(w - r * w), int(h - b * h)))


def _green_bbox(im, pad=4):
    """CAN 컨버터 사진에서 초록 기판 영역만 남긴다 (둘레의 한글 주석 제거)."""
    px = im.load()
    w, h = im.size
    xs, ys = [], []
    for y in range(h):
        for x in range(w):
            r, g, b = px[x, y]
            if g > 90 and g > r + 25 and g > b + 25:
                xs.append(x)
                ys.append(y)
    return im.crop((max(0, min(xs) - pad), max(0, min(ys) - pad), min(w, max(xs) + pad), min(h, max(ys) + pad)))


def make_fig2():
    W, H = 7.16, 2.45
    prs, slide = new_pres(W, H)
    t = TC(slide)
    tmp = os.path.join(PH, 'crops')
    os.makedirs(tmp, exist_ok=True)

    def pic(im, fname, x, y, w=None, h=None, name=None):
        path = os.path.join(tmp, fname)
        im.save(path, quality=92)
        if w is not None and h is None:
            h = w * im.height / im.width
        elif h is not None and w is None:
            w = h * im.width / im.height
        p = slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
        if name:
            p.name = name
        return x, y, w, h

    # ---------- (a) 선박 ----------
    im = Image.open(os.path.join(PH, 'vessel_side.jpg')).crop((140, 300, 1340, 1040))
    wa = 2.20
    ha = wa * im.height / im.width
    pic(im, 'a.jpg', 0.02, 0.05 + (2.05 - ha) / 2, w=wa, name='photo_a_vessel')

    # ---------- (b) 선미 기관실 ----------
    x0, y0, x1, y1 = 300, 360, 760, 960
    imb = Image.open(os.path.join(PH, 'stern_bay.jpg')).crop((x0, y0, x1, y1))
    wb = 1.40
    hb = wb * imb.height / imb.width
    xb, yb = 2.32, 0.05 + (2.05 - hb) / 2
    pic(imb, 'b.jpg', xb, yb, w=wb, name='photo_b_stern_bay')

    def bp(ix, iy):
        return xb + ix / imb.width * wb, yb + iy / imb.height * hb

    targets = {'shift': (638, 691), 'steer': (413, 539), 'pivot': (485, 858)}
    labels = [('steer', ['Steering servo', 'cylinder'], (120, 40)),
              ('shift', ['Shift servo', 'cylinder'], (345, 150)),
              ('pivot', ['Drive steering', 'pivot'], (150, 560))]
    lw_, lh_ = 0.66, 0.22
    for key, lines, (lx, ly) in labels:
        tx, ty = targets[key]
        pt = bp(tx - x0, ty - y0)
        pl = bp(lx, ly)
        start = (pl[0], pl[1] + lh_ / 2) if pt[1] > pl[1] else (pl[0], pl[1] - lh_ / 2)
        t.arrow(start, pt, color=C_YEL, lw=1.0)
        t.node(pl[0] - lw_ / 2, pl[1] - lh_ / 2, lw_, lh_, [T(s) for s in lines], C_PLANT, size=6.0, lw=0.6)

    # ---------- (c) 제어 하드웨어 구성 (원본 5번 슬라이드 형식, Pixhawk 제외) ----------
    cx0 = 3.84
    # 선박 제어부 점선
    t.rect(cx0 + 0.02, 0.58, 3.26, 1.50, C_BOUND, None, lw=1.2, radius=0.18, dash=True, name='boat_control_boundary')
    t.label(cx0 + 1.70, 0.43, 1.58, 0.13, [T('Boat control hardware')], size=6.4, color=C_BOUND,
            bold=True, align='right')

    # Control PC
    pc = _media('image12.jpg')
    px_, py_, pw_, ph_ = pic(pc, 'pc.jpg', cx0 + 0.10, 0.06, w=0.62, name='control_pc')
    t.label(cx0 + 0.78, 0.08, 0.9, 0.34, [T(('Control PC', 'b')), T('ROS 2')], size=6.6, align='left')

    # PLC (빨간 테두리)
    plc = _media('image15.jpeg')
    lx_, ly_, lw2, lh2 = pic(plc, 'plc.jpg', cx0 + 0.14, 0.86, w=0.58, name='plc')
    t.rect(lx_, ly_, lw2, lh2, C_BOUND, None, lw=1.0, radius=0.01)
    t.label(lx_ - 0.10, ly_ + lh2 + 0.02, lw2 + 0.20, 0.13, [T(('LS XGB PLC', 'b'))], size=6.4)

    # Control PC → PLC
    xl = px_ + pw_ / 2
    t.arrow((xl, py_ + ph_), (xl, ly_), color=C_LINK, lw=1.3, head=False)
    t.label(xl + 0.04, 0.43, 0.6, 0.13, [T('Ethernet')], size=6.0, color=C_PROTO, align='left')

    # 버스
    xbus = cx0 + 1.02
    t.arrow((lx_ + lw2, ly_ + lh2 / 2), (xbus, ly_ + lh2 / 2), color=C_LINK, lw=1.3, head=False)

    # CAN 컨버터
    can = _green_bbox(_media('image16.png'))
    cx_, cy_, cw_, ch_ = pic(can, 'can.jpg', cx0 + 1.18, 0.66, w=0.50, name='can_converter')
    t.label(cx_ - 0.1, cy_ + ch_ + 0.02, cw_ + 0.2, 0.13, [T(('CAN converter', 'b'))], size=6.4)
    yc = cy_ + ch_ / 2
    # 서보 실린더
    srv = _crop_frac(_media('image17.jpeg'), 0.0645, 0.1363, 0.1306, 0.1351)
    sx_, sy_, sw_, sh_ = pic(srv, 'servo.jpg', cx0 + 1.18, 1.42, w=0.84, name='servo_cylinder')
    t.label(sx_ - 0.1, sy_ + sh_ + 0.02, sw_ + 0.2, 0.13, [T(('Servo cylinders', 'b'))], size=6.4)
    ys = sy_ + sh_ / 2
    t.arrow((xbus, yc), (xbus, ys), color=C_LINK, lw=1.3, head=False)
    t.arrow((xbus, yc), (cx_, yc), color=C_LINK, lw=1.3)
    t.arrow((xbus, ys), (sx_, ys), color=C_LINK, lw=1.3)
    t.label(xbus - 0.31, yc - 0.15, 0.30, 0.13, [T('RS232C')], size=6.0, color=C_PROTO, align='right')
    t.label(xbus - 0.53, ys + 0.02, 0.50, 0.13, [T('EtherCAT')], size=6.0, color=C_PROTO, align='right')

    # 엔진과 선미 드라이브
    eng = _crop_frac(_media('image14.jpg'), 0.19817, 0.17253, 0.13359, 0.46496)
    ex_, ey_, ew_, eh_ = pic(eng, 'engine.jpg', cx0 + 2.32, 0.88, w=0.90, name='engine_sterndrive')
    # CAN 컨버터 → 엔진 회전수
    xe1 = ex_ + ew_ * 0.35
    t.poly([(cx_ + cw_, yc), (xe1, yc), (xe1, ey_)], color=C_LINK, lw=1.3)
    t.label(cx_ + cw_ + 0.03, yc - 0.15, 0.4, 0.13, [T('CAN')], size=6.0, color=C_PROTO, align='left')
    t.label(xe1 - 0.75, ey_ - 0.27, 0.72, 0.13, [T(('Engine speed', 'b'))], size=6.4, align='right')
    # 서보 실린더 → 시프트와 조향
    xe2 = ex_ + ew_ * 0.80
    t.poly([(sx_ + sw_, ys + 0.10), (xe2, ys + 0.10), (xe2, ey_ + eh_)], color=C_LINK, lw=1.3)
    t.label(sx_ + sw_ + 0.05, ys + 0.13, 1.25, 0.13, [T(('Shift and steering', 'b'))], size=6.4, align='left')

    # 패널 이름
    for xx, s in [(0.02 + wa / 2, '(a)'), (xb + wb / 2, '(b)'), (cx0 + 1.65, '(c)')]:
        t.label(xx - 0.2, 2.28, 0.4, 0.14, [T(s)], size=7.5)

    out = os.path.join(OUTDIR, 'fig2_vessel.pptx')
    prs.save(out)
    return out


if __name__ == '__main__':
    print(make_fig1())
    print(make_fig2())
